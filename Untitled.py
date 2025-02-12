#!/usr/bin/env python
# coding: utf-8

# In[ ]:


import requests
import pandas as pd
import time
from datetime import datetime
import openpyxl
from openpyxl.styles import PatternFill, Font
import os

class CryptoAnalyzer:
    def __init__(self):
        self.api_url = "https://api.coingecko.com/api/v3"
        self.excel_file = "crypto_data_live.xlsx"
        
    def fetch_top_50_crypto(self):
        
        try:
            endpoint = f"{self.api_url}/coins/markets"
            params = {
                'vs_currency': 'usd',
                'order': 'market_cap_desc',
                'per_page': 50,
                'page': 1,
                'sparkline': False
            }
            
            response = requests.get(endpoint, params=params)
            response.raise_for_status()
            return response.json()
            
        except requests.RequestException as e:
            print(f"Error fetching data: {e}")
            return None

    def process_data(self, data):
        
        if not data:
            return None
            
        df = pd.DataFrame(data)
        df = df[[
            'name', 
            'symbol', 
            'current_price', 
            'market_cap', 
            'total_volume',
            'price_change_percentage_24h'
        ]]
        
        df.columns = [
            'Cryptocurrency Name',
            'Symbol',
            'Current Price (USD)',
            'Market Capitalization',
            '24h Trading Volume',
            'Price Change 24h (%)'
        ]
        
        return df

    def analyze_data(self, df):
        
        if df is None:
            return None
            
        analysis = {
            'Top 5 by Market Cap': df.head(),
            'Average Price': df['Current Price (USD)'].mean(),
            'Highest 24h Change': df.nlargest(1, 'Price Change 24h (%)'),
            'Lowest 24h Change': df.nsmallest(1, 'Price Change 24h (%)')
        }
        
        return analysis

    def update_excel(self, df, analysis):
        
        if df is None or analysis is None:
            return
            
        
        with pd.ExcelWriter(self.excel_file, engine='openpyxl') as writer:
            
            df.to_excel(writer, sheet_name='Live Data', index=False)
            
            
            analysis['Top 5 by Market Cap'].to_excel(writer, sheet_name='Analysis', index=False)
            
            
            summary_df = pd.DataFrame({
                'Metric': ['Average Price (USD)', 
                          'Highest 24h Change (%)', 
                          'Lowest 24h Change (%)'],
                'Value': [
                    f"${analysis['Average Price']:.2f}",
                    f"{analysis['Highest 24h Change']['Price Change 24h (%)'].iloc[0]:.2f}%",
                    f"{analysis['Lowest 24h Change']['Price Change 24h (%)'].iloc[0]:.2f}%"
                ]
            })
            
            summary_df.to_excel(writer, sheet_name='Analysis', 
                              startrow=len(analysis['Top 5 by Market Cap']) + 2, 
                              index=False)

    def format_excel(self):
        
        wb = openpyxl.load_workbook(self.excel_file)
        
        # Format Live Data sheet
        ws_data = wb['Live Data']
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True)
        
        for cell in ws_data[1]:
            cell.fill = header_fill
            cell.font = header_font
            
        # Format Analysis sheet
        ws_analysis = wb['Analysis']
        for cell in ws_analysis[1]:
            cell.fill = header_fill
            cell.font = header_font
            
        wb.save(self.excel_file)

    def run_live_update(self, interval_minutes=5):
        
        while True:
            print(f"\nFetching data at {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            
            
            raw_data = self.fetch_top_50_crypto()
            df = self.process_data(raw_data)
            analysis = self.analyze_data(df)
            
           
            self.update_excel(df, analysis)
            self.format_excel()
            
            print("Excel file updated successfully!")
            print(f"Next update in {interval_minutes} minutes...")
            
            time.sleep(interval_minutes * 60)

def main():
    analyzer = CryptoAnalyzer()
    try:
        analyzer.run_live_update()
    except KeyboardInterrupt:
        print("\nStopping the live update process...")

if __name__ == "__main__":
    main()


# In[ ]:




